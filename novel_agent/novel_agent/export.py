"""导出：整书 Markdown / TXT / DOCX。"""
from __future__ import annotations

from pathlib import Path

from .memory import StoryBible


def render_markdown(bible: StoryBible) -> str:
    d = bible.data
    lines = [f"# {d['title']}", "", f"> {d.get('logline','')}", ""]
    if d.get("world_setting"):
        lines += ["## 世界观", "", d["world_setting"], ""]
    if d.get("characters"):
        lines += ["## 主要人物", ""]
        for c in d["characters"]:
            lines.append(f"- **{c.get('name')}**（{c.get('role')}）：{c.get('goal')}；{c.get('traits')}")
        lines.append("")
    for rec in d["chapters"]:
        lines += [f"## 第{rec['index']}章 {rec.get('title','')}", "", rec.get("text", ""), ""]
    return "\n".join(lines)


def export_all(bible: StoryBible, out_dir: Path) -> dict[str, Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    md = render_markdown(bible)
    md_path = out_dir / "full_novel.md"
    md_path.write_text(md, encoding="utf-8")

    txt_path = out_dir / "full_novel.txt"
    txt_path.write_text(md.replace("#", "").replace("**", ""), encoding="utf-8")

    docx_path = out_dir / "full_novel.docx"
    try:
        from docx import Document

        doc = Document()
        d = bible.data
        doc.add_heading(d["title"], level=0)
        for rec in d["chapters"]:
            doc.add_heading(f"第{rec['index']}章 {rec.get('title','')}", level=1)
            for para in rec.get("text", "").split("\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
        doc.save(str(docx_path))
    except Exception as e:  # python-docx 缺失时不阻断主流程
        docx_path = None  # type: ignore[assignment]
    return {"md": md_path, "txt": txt_path, "docx": docx_path}  # type: ignore[dict-item]
