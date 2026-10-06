"""全書导出：Markdown 以外的 PDF（排版美化）与 Word（docx）格式。"""
from __future__ import annotations

from datetime import datetime
from typing import Iterable


def _split_paragraphs(text: str) -> list[str]:
    """按空行/换行切分段落，去掉空白段。"""
    paras = [p.strip() for p in text.replace("\r\n", "\n").split("\n")]
    out: list[str] = []
    buf = ""
    for p in paras:
        if p:
            buf = f"{buf}\n{p}" if buf else p
        elif buf:
            out.append(buf)
            buf = ""
    if buf:
        out.append(buf)
    return out


def _meta(chapters: list[dict]) -> dict:
    total_words = sum(len(c["text"] or "") for c in chapters)
    return {
        "chapter_count": len(chapters),
        "total_words": total_words,
        "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


# --------------------------------------------------------------------- #
# PDF：reportlab 排版，内置 STSong-Light CID 字体，无需字体文件
# --------------------------------------------------------------------- #
def build_pdf(book_name: str, chapters: list[dict], author: str = "", bio: str = "") -> bytes:
    from io import BytesIO

    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm, mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas as _canvas
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        HRFlowable,
    )
    from reportlab.lib.styles import ParagraphStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    FONT = "STSong-Light"
    INK = HexColor("#1c1917")        # 正文墨色
    GOLD = HexColor("#8a6d3b")       # 点缀金
    GREY = HexColor("#78716c")

    meta = _meta(chapters)
    buf = BytesIO()

    class NumberedCanvas(_canvas.Canvas):
        """页脚：居中页码 + 右上角书名，封面页除外。"""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states: list[dict] = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved_page_states)
            for i, state in enumerate(self._saved_page_states):
                self.__dict__.update(state)
                if i > 0:  # 封面不标页码
                    self._draw_footer(i + 1)
                _canvas.Canvas.showPage(self)
            _canvas.Canvas.save(self)

        def _draw_footer(self, page_num: int):
            width, height = A4
            self.setFont(FONT, 8)
            self.setFillColor(GREY)
            self.drawCentredString(width / 2, 12 * mm, f"— {page_num} —")
            self.drawRightString(width - 2 * cm, height - 12 * mm, book_name[:20])

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2.4 * cm,
        rightMargin=2.4 * cm,
        topMargin=2.4 * cm,
        bottomMargin=2.2 * cm,
        title=book_name,
        author=author or "墨卷 NovelForge",
    )

    title_style = ParagraphStyle(
        "booktitle", fontName=FONT, fontSize=30, leading=44,
        textColor=INK, alignment=1, spaceAfter=6 * mm,
    )
    sub_style = ParagraphStyle(
        "subtitle", fontName=FONT, fontSize=12, leading=20,
        textColor=GREY, alignment=1,
    )
    chapter_style = ParagraphStyle(
        "chapter", fontName=FONT, fontSize=17, leading=26,
        textColor=INK, spaceBefore=4 * mm, spaceAfter=2 * mm,
    )
    bio_head_style = ParagraphStyle(
        "biohead", fontName=FONT, fontSize=15, leading=24,
        textColor=INK, spaceAfter=4 * mm,
    )
    body_style = ParagraphStyle(
        "body", fontName=FONT, fontSize=11.5, leading=21,
        textColor=INK, firstLineIndent=23, spaceAfter=2.2 * mm,
        alignment=4,  # justify
    )

    story: list = [
        Spacer(1, 5.5 * cm),
        Paragraph(book_name, title_style),
        HRFlowable(width="38%", thickness=1, color=GOLD, hAlign="CENTER", spaceAfter=6 * mm),
    ]
    if author:
        story.append(Paragraph(author, ParagraphStyle(
            "author", fontName=FONT, fontSize=14, leading=24,
            textColor=INK, alignment=1, spaceAfter=4 * mm,
        )))
    story += [
        Paragraph(
            f"全书共 {meta['chapter_count']} 章 · 约 {meta['total_words']} 字",
            sub_style,
        ),
        Paragraph(f"导出于 {meta['exported_at']}", sub_style),
        Spacer(1, 4 * cm),
        Paragraph("由 墨卷 NovelForge · AI 长篇小说生成平台 排版导出", sub_style),
        PageBreak(),
    ]

    if bio:
        story += [
            Spacer(1, 1.5 * cm),
            Paragraph("内容简介", bio_head_style),
            HRFlowable(width="100%", thickness=0.6, color=HexColor("#d6d3d1"),
                       spaceBefore=1 * mm, spaceAfter=4 * mm),
        ]
        for para in _split_paragraphs(bio):
            safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe.replace("\n", "<br/>"), body_style))
        story += [
            Spacer(1, 1.2 * cm),
            Paragraph(
                f"© {datetime.now().year} {author or '本书作者'}。本书内容由 AI 辅助生成，版权归作者所有，转载请注明出处。",
                ParagraphStyle("copyright", fontName=FONT, fontSize=9, leading=16,
                               textColor=GREY),
            ),
            PageBreak(),
        ]

    for c in chapters:
        story.append(Paragraph(f"第{c['idx']}章　{c['title']}", chapter_style))
        story.append(
            HRFlowable(width="100%", thickness=0.6, color=HexColor("#d6d3d1"),
                       spaceBefore=1 * mm, spaceAfter=4 * mm)
        )
        for para in _split_paragraphs(c["text"] or ""):
            safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe.replace("\n", "<br/>"), body_style))
        story.append(Spacer(1, 6 * mm))

    doc.build(story, canvasmaker=NumberedCanvas)
    return buf.getvalue()


# --------------------------------------------------------------------- #
# Word：python-docx，正文宋体小四、1.5 倍行距、首行缩进两字符
# --------------------------------------------------------------------- #
def build_docx(book_name: str, chapters: list[dict], author: str = "", bio: str = "") -> bytes:
    from io import BytesIO

    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    meta = _meta(chapters)
    d = docx.Document()

    # 默认样式：中文宋体、西文 Times New Roman、小四
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x1C, 0x19, 0x17)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def set_ea(run, font="宋体"):
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)

    # 封面式标题块
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(book_name)
    r.font.size = Pt(26)
    r.font.bold = True
    set_ea(r, "黑体")
    t.paragraph_format.space_before = Pt(120)

    if author:
        ap = d.add_paragraph()
        ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ar = ap.add_run(author)
        ar.font.size = Pt(14)
        set_ea(ar, "楷体")
        ap.paragraph_format.space_after = Pt(18)

    sub = d.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        f"全书共 {meta['chapter_count']} 章 · 约 {meta['total_words']} 字\n"
        f"导出于 {meta['exported_at']} · 墨卷 NovelForge"
    )
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x78, 0x71, 0x6C)
    set_ea(r)

    if bio:
        d.add_page_break()
        bh = d.add_paragraph()
        br = bh.add_run("内容简介")
        br.font.size = Pt(15)
        br.font.bold = True
        set_ea(br, "黑体")
        bh.paragraph_format.space_after = Pt(10)
        for para in _split_paragraphs(bio):
            p = d.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.first_line_indent = Pt(24)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(para)
            r.font.size = Pt(12)
            set_ea(r)
        cp = d.add_paragraph()
        cp.paragraph_format.space_before = Pt(24)
        cr = cp.add_run(
            f"© {datetime.now().year} {author or '本书作者'}。"
            "本书内容由 AI 辅助生成，版权归作者所有，转载请注明出处。"
        )
        cr.font.size = Pt(9)
        cr.font.color.rgb = RGBColor(0x78, 0x71, 0x6C)
        set_ea(cr)

    for c in chapters:
        d.add_page_break()
        h = d.add_paragraph()
        r = h.add_run(f"第{c['idx']}章　{c['title']}")
        r.font.size = Pt(16)
        r.font.bold = True
        set_ea(r, "黑体")
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(10)

        for para in _split_paragraphs(c["text"] or ""):
            p = d.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.first_line_indent = Pt(24)  # 首行缩进两字符
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(para)
            r.font.size = Pt(12)
            set_ea(r)

    out = BytesIO()
    d.save(out)
    return out.getvalue()
